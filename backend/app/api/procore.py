"""Procore submittal review: reviewer comments, the Comment Log Sheet, revisions.

The Procore round-trip is a document, not a chat: the reviewers (HSE, FLP/Fire
Protection, FMOM/Building Management, Building Operations, CTPS, the Proponent
PI ...) send comments on Rev 0, IHP answers every one of them, the log is
printed as PR-XXXXX Comments Logs, and the answered package goes back as
Rev 1. So this module owns:

* the comment log rows (reviewer, department, discipline, comment, response,
  action taken, open/closed, the revision it was raised on),
* the manual Procore status on the project (Under Review | Revise and
  Resubmit | Approved as Noted | Approved), submittal date and revision,
* the stamped Procore Submittal PDF dropzone,
* the Comment Log Sheet export (.xlsx and .pdf), and
* the "Generate Revision N Package" step that bumps the revision.
"""

from datetime import date, datetime, timezone
from pathlib import Path
from tempfile import NamedTemporaryFile

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import FileResponse
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..core.rbac import CAP_SOW_MANAGE, get_current_user, require_capability
from ..db import get_db
from ..models import Attachment, ProcoreComment, Project, User
from ..services import storage, workflow
from .projects import get_project_or_404

router = APIRouter(prefix="/projects/{project_id}/procore", tags=["procore"])

#: The reviewers the team actually sends submittals to, in the order they appear
#: on the comment log. Free text is always allowed as well.
REVIEWERS = [
    "HSE Team",
    "FLP / Fire Protection",
    "FMOM / Building Management",
    "Building Operations",
    "CTPS",
    "Proponent / PI",
]

#: The Procore review outcome, chosen manually in the app.
PROCORE_STATUSES = [
    "Under Review",
    "Revise and Resubmit",
    "Approved as Noted",
    "Approved",
]

APPROVED_STATUSES = ("Approved", "Approved as Noted")

COMMENT_STATUSES = ("open", "closed")


# ---------- request bodies ----------
class CommentIn(BaseModel):
    reviewer: str = ""
    department: str = ""
    discipline: str = ""
    comment: str = ""
    ihp_response: str | None = None
    action_taken: str | None = None
    status: str = "open"
    revision: int | None = None


class CommentPatch(BaseModel):
    reviewer: str | None = None
    department: str | None = None
    discipline: str | None = None
    comment: str | None = None
    ihp_response: str | None = None
    action_taken: str | None = None
    status: str | None = None
    revision: int | None = None
    item_no: int | None = None


class SubmittalPatch(BaseModel):
    """The manual Procore review state, as the reviewer sets it in Procore."""

    procore_status: str | None = None
    procore_submittal_date: date | None = None
    procore_revision: int | None = None


def _out(row: ProcoreComment) -> dict:
    return {
        "id": row.id,
        "item_no": row.item_no,
        "reviewer": row.reviewer or "",
        "department": row.department or "",
        "discipline": row.discipline or "",
        "comment": row.comment or "",
        "ihp_response": row.ihp_response or "",
        "action_taken": row.action_taken or "",
        "status": row.status or "open",
        "revision": row.revision or 0,
        "created_at": row.created_at.isoformat() if row.created_at else None,
    }


def _comments(db: Session, project_id: int) -> list[ProcoreComment]:
    return list(
        db.scalars(
            select(ProcoreComment)
            .where(ProcoreComment.project_id == project_id)
            .order_by(ProcoreComment.item_no, ProcoreComment.id)
        ).all()
    )


def _next_item_no(db: Session, project_id: int) -> int:
    rows = _comments(db, project_id)
    return max((row.item_no or 0) for row in rows) + 1 if rows else 1


@router.get("/reviewers")
def reviewer_list(_user: User = Depends(get_current_user)):
    """The quick-select reviewer list plus the status vocabulary."""
    return {"reviewers": REVIEWERS, "statuses": PROCORE_STATUSES}


@router.get("")
def procore_state(
    project_id: int,
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
):
    """Submittal state + the whole comment log for the panel."""
    project = get_project_or_404(db, project_id)
    rows = _comments(db, project.id)
    return {
        "procore_status": project.procore_status,
        "procore_submittal_date": (
            project.procore_submittal_date.isoformat()
            if project.procore_submittal_date
            else None
        ),
        "revision": project.procore_revision or 0,
        "reviewers": REVIEWERS,
        "statuses": PROCORE_STATUSES,
        "comments": [_out(row) for row in rows],
        "open_comments": sum(1 for row in rows if (row.status or "open") == "open"),
    }


@router.patch("")
def set_procore_state(
    project_id: int,
    body: SubmittalPatch,
    db: Session = Depends(get_db),
    user: User = Depends(require_capability(CAP_SOW_MANAGE)),
):
    """Set the Procore status / submittal date / current revision by hand."""
    project = get_project_or_404(db, project_id)
    if body.procore_status is not None:
        if body.procore_status and body.procore_status not in PROCORE_STATUSES:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Unknown Procore status: {body.procore_status}",
            )
        project.procore_status = body.procore_status or None
    if body.procore_submittal_date is not None:
        project.procore_submittal_date = body.procore_submittal_date
    if body.procore_revision is not None:
        project.procore_revision = max(0, body.procore_revision)
    workflow.log_action(
        db,
        user,
        "procore:status",
        project,
        {
            "status": project.procore_status,
            "submittal_date": (
                project.procore_submittal_date.isoformat()
                if project.procore_submittal_date
                else None
            ),
            "revision": project.procore_revision,
        },
    )
    db.commit()
    return {
        "procore_status": project.procore_status,
        "procore_submittal_date": (
            project.procore_submittal_date.isoformat()
            if project.procore_submittal_date
            else None
        ),
        "revision": project.procore_revision or 0,
    }


# ---------- comment log rows ----------
@router.post("/comments", status_code=status.HTTP_201_CREATED)
def add_comment(
    project_id: int,
    body: CommentIn,
    db: Session = Depends(get_db),
    user: User = Depends(require_capability(CAP_SOW_MANAGE)),
):
    project = get_project_or_404(db, project_id)
    if not (body.comment or "").strip():
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="The comment text is empty.",
        )
    row = ProcoreComment(
        project_id=project.id,
        item_no=_next_item_no(db, project.id),
        reviewer=body.reviewer.strip(),
        department=body.department.strip(),
        discipline=body.discipline.strip(),
        comment=body.comment.strip(),
        ihp_response=(body.ihp_response or "").strip() or None,
        action_taken=(body.action_taken or "").strip() or None,
        status=body.status if body.status in COMMENT_STATUSES else "open",
        revision=(
            body.revision
            if body.revision is not None
            else (project.procore_revision or 0)
        ),
        created_by_id=user.id,
    )
    db.add(row)
    db.flush()
    workflow.log_action(
        db,
        user,
        "procore:comment",
        project,
        {"item_no": row.item_no, "reviewer": row.reviewer, "comment": row.comment[:200]},
    )
    db.commit()
    db.refresh(row)
    return _out(row)


@router.patch("/comments/{comment_id}")
def update_comment(
    project_id: int,
    comment_id: int,
    body: CommentPatch,
    db: Session = Depends(get_db),
    user: User = Depends(require_capability(CAP_SOW_MANAGE)),
):
    project = get_project_or_404(db, project_id)
    row = db.get(ProcoreComment, comment_id)
    if row is None or row.project_id != project.id:
        raise HTTPException(status_code=404, detail="Comment not found")
    data = body.model_dump(exclude_unset=True)
    if "status" in data and data["status"] not in COMMENT_STATUSES:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="status must be open or closed",
        )
    for key, value in data.items():
        if isinstance(value, str):
            value = value.strip()
        if key in {"ihp_response", "action_taken"} and value == "":
            value = None
        setattr(row, key, value)
    db.flush()
    workflow.log_action(
        db, user, "procore:comment_update", project, {"item_no": row.item_no}
    )
    db.commit()
    db.refresh(row)
    return _out(row)


@router.delete("/comments/{comment_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_comment(
    project_id: int,
    comment_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_capability(CAP_SOW_MANAGE)),
):
    project = get_project_or_404(db, project_id)
    row = db.get(ProcoreComment, comment_id)
    if row is None or row.project_id != project.id:
        raise HTTPException(status_code=404, detail="Comment not found")
    db.delete(row)
    workflow.log_action(
        db, user, "procore:comment_delete", project, {"item_no": row.item_no}
    )
    db.commit()
    return None


# ---------- Comment Log Sheet ----------
_COMMENT_HEADERS = [
    "Item #",
    "Reviewer",
    "Department / Discipline",
    f"Comment",
    "IHP Response",
    "Action Taken",
    "Status",
]
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def comment_log_xlsx(project: Project, comments: list[ProcoreComment], target: Path) -> Path:
    """Write PR-XXXXX Comments Logs as an .xlsx sheet with the title block."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Comments Log"

    ws["A1"] = f"{project.pr_number} - Procore Comments Log"
    ws["A1"].font = Font(size=14, bold=True)
    ws.merge_cells("A1:G1")
    revision = project.procore_revision or 0
    meta = [
        ("Project", f"{project.pr_number}  {project.title or ''}".strip()),
        ("Proponent / PI", project.pi_name or ""),
        ("Submittal revision", f"Rev {revision}"),
        (
            "Submittal date",
            project.procore_submittal_date.strftime("%d %B %Y")
            if project.procore_submittal_date
            else "",
        ),
        ("Procore status", project.procore_status or ""),
        (
            "Printed",
            datetime.now(timezone.utc).strftime("%d %B %Y %H:%M UTC"),
        ),
    ]
    for index, (label, value) in enumerate(meta, start=3):
        ws.cell(row=index, column=1, value=label).font = Font(bold=True)
        ws.cell(row=index, column=2, value=value)
        ws.merge_cells(start_row=index, start_column=2, end_row=index, end_column=7)

    head_row = 3 + len(meta) + 1
    for column, title in enumerate(_COMMENT_HEADERS, start=1):
        cell = ws.cell(row=head_row, column=column, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color="1F3864")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = _BORDER

    row_index = head_row
    for comment in comments:
        row_index += 1
        dept = " / ".join(
            part for part in [comment.department or "", comment.discipline or ""] if part
        )
        values = [
            comment.item_no,
            comment.reviewer or "",
            dept,
            comment.comment or "",
            comment.ihp_response or "",
            comment.action_taken or "",
            (comment.status or "open").title(),
        ]
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row_index, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = _BORDER

    if not comments:
        ws.cell(
            row=head_row + 1,
            column=1,
            value="No reviewer comments recorded yet.",
        )

    for column, width in zip("ABCDEFG", (8, 22, 24, 62, 52, 34, 10)):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    wb.save(target)
    return target


def _soffice_pdf(source: Path) -> Path | None:
    from pathlib import Path as _Path

    from .sow_boq import _soffice_pdf as convert

    return convert(_Path(source))


@router.get("/comments/log.xlsx")
def comment_log_excel(
    project_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """Download the comment log as the team's .xlsx sheet."""
    project = get_project_or_404(db, project_id)
    comments = _comments(db, project.id)
    out_dir = storage.project_dir(project.pr_number, "generated")
    out_dir.mkdir(parents=True, exist_ok=True)
    target = out_dir / f"{project.pr_number} Comments Logs.xlsx"
    comment_log_xlsx(project, comments, target)
    workflow.log_action(
        db, user, "procore:log_xlsx", project, {"rows": len(comments)}
    )
    db.commit()
    return FileResponse(
        target,
        filename=target.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@router.get("/comments/log.pdf")
def comment_log_pdf(
    project_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """Download the comment log as a PDF (LibreOffice print of the sheet)."""
    project = get_project_or_404(db, project_id)
    comments = _comments(db, project.id)
    out_dir = storage.project_dir(project.pr_number, "generated")
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx = out_dir / f"{project.pr_number} Comments Logs.xlsx"
    comment_log_xlsx(project, comments, xlsx)
    pdf = _soffice_pdf(xlsx)
    if pdf is None or not Path(pdf).exists():
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="The PDF converter (LibreOffice) is unavailable right now.",
        )
    target = out_dir / f"{project.pr_number} Comments Logs.pdf"
    Path(pdf).replace(target)
    workflow.log_action(db, user, "procore:log_pdf", project, {"rows": len(comments)})
    db.commit()
    return FileResponse(target, filename=target.name, media_type="application/pdf")


# ---------- stamped Procore submittal PDF ----------
@router.post("/submittal", status_code=status.HTTP_201_CREATED)
async def upload_submittal(
    project_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(require_capability(CAP_SOW_MANAGE)),
):
    """Dropzone for the stamped Procore Submittal PDF that came back reviewed."""
    project = get_project_or_404(db, project_id)
    filename = file.filename or "Procore Submittal.pdf"
    if not filename.lower().endswith(".pdf"):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="The reviewed submittal must be a PDF.",
        )
    data = await file.read()
    if not data:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="The file is empty."
        )
    stored_path, version = storage.store_upload(
        project.pr_number, "DELIVERABLE", filename, data
    )
    attachment = Attachment(
        project_id=project.id,
        stage="DELIVERABLE",
        filename=filename,
        stored_path=str(stored_path),
        content_type="application/pdf",
        size_bytes=len(data),
        version=version,
        uploaded_by_id=user.id,
    )
    db.add(attachment)
    db.flush()
    workflow.log_action(
        db,
        user,
        "procore:submittal_pdf",
        project,
        {"file": filename, "version": version, "bytes": len(data)},
    )
    db.commit()
    return {
        "id": attachment.id,
        "filename": filename,
        "version": version,
        "size_bytes": len(data),
    }


# ---------- Generate Revision N Package ----------
@router.post("/revise")
def generate_revision(
    project_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_capability(CAP_SOW_MANAGE)),
):
    """Bump the submittal revision (Rev 0 -> Rev 1) for the next binder.

    The binder itself is produced by POST /package/pdf, which picks the current
    revision up from the project - so this step records the decision, keeps the
    previous revision's comments in the log (they stay attached to the revision
    they were raised on) and returns what the new package must still answer.
    """
    project = get_project_or_404(db, project_id)
    previous = project.procore_revision or 0
    project.procore_revision = previous + 1
    project.procore_status = "Under Review"
    project.procore_submittal_date = datetime.now(timezone.utc).date()
    rows = _comments(db, project.id)
    open_rows = [row for row in rows if (row.status or "open") == "open"]
    workflow.log_action(
        db,
        user,
        "procore:revise",
        project,
        {"from": previous, "to": project.procore_revision, "open": len(open_rows)},
    )
    db.commit()
    return {
        "revision": project.procore_revision,
        "previous_revision": previous,
        "procore_status": project.procore_status,
        "procore_submittal_date": project.procore_submittal_date.isoformat(),
        "comments": len(rows),
        "open_comments": len(open_rows),
        "must_answer": [row.item_no for row in open_rows],
    }
