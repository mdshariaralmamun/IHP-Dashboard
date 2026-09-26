"""Stage 4 SOW (Scope of Work) and BOQ/MTO (Bill of Quantities / Material Take-Off) endpoints.

Features:
- SOW revision tracking (Rev-0, Rev-1, Rev-2...) with Procore review comments
- Multi-trade BOQ/MTO line-item manager
- Excel (.xlsx) export using openpyxl
"""

from datetime import datetime, timezone
import io
from pathlib import Path
from tempfile import NamedTemporaryFile

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import FileResponse, StreamingResponse
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..core.rbac import (
    CAP_BOQ_MANAGE,
    CAP_SOW_MANAGE,
    effective_permissions,
    get_current_user,
    require_capability,
)
from ..db import get_db
from ..models import BoqMtoItem, Project, SowRecord, User
from ..schemas import (
    BoqItemCreate,
    BoqItemOut,
    BoqItemUpdate,
    SowOut,
    SowRevisionCreate,
)
from ..services import sow_boq_mto_docgen as docgen
from ..services import storage, workflow
from .projects import get_project_or_404

router = APIRouter(prefix="/projects/{project_id}", tags=["sow_boq"])


# ---------- SOW Endpoints ----------
@router.get("/sow", response_model=list[SowOut])
def list_sow_revisions(
    project_id: int,
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
):
    """List all SOW revisions for a project."""
    project = get_project_or_404(project_id, db)
    return project.sow_records


@router.post("/sow", response_model=SowOut)
def create_sow_revision(
    project_id: int,
    body: SowRevisionCreate,
    db: Session = Depends(get_db),
    user: User = Depends(require_capability(CAP_SOW_MANAGE)),
):
    """Create a new SOW revision (Rev-0, Rev-1, Rev-2) with Procore feedback."""
    project = get_project_or_404(project_id, db)

    # ICR fast-track skips the SOW step entirely; only MTO and beyond apply.
    if project.disposition == "ICR":
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="SOW is not part of the ICR fast-track workflow; this project routes to MTO only.",
        )

    # Determine revision name if not specified
    existing_count = len(project.sow_records)
    rev_name = body.revision_name or f"Rev-{existing_count}"

    sow = SowRecord(
        project_id=project.id,
        revision_name=rev_name,
        status="draft",
        scope_text=body.scope_text or f"Detailed Scope of Work for {project.title}",
        trade_sections=body.trade_sections,
        procore_comments=body.procore_comments,
        created_by_id=user.id,
    )
    db.add(sow)

    if project.stage in (workflow.EAR_APPROVED, workflow.DISPOSITION):
        project.stage = workflow.SOW_DRAFT

    workflow.log_action(
        db,
        user,
        f"sow:created:{rev_name}",
        project,
        {"revision": rev_name, "procore_comments": bool(body.procore_comments)},
    )

    db.commit()
    db.refresh(sow)
    return sow


@router.patch("/sow/{sow_id}", response_model=SowOut)
def update_sow_status(
    project_id: int,
    sow_id: int,
    status_val: str = Query(..., pattern="^(draft|in_review|approved)$"),
    db: Session = Depends(get_db),
    user: User = Depends(require_capability(CAP_SOW_MANAGE)),
):
    """Update SOW revision status (draft -> in_review -> approved)."""
    project = get_project_or_404(project_id, db)
    sow = db.get(SowRecord, sow_id)
    if not sow or sow.project_id != project.id:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="SOW revision not found"
        )

    sow.status = status_val
    sow.updated_at = datetime.now(timezone.utc)

    if status_val == "approved":
        project.stage = workflow.SOW_APPROVED
        workflow.log_action(db, user, "stage:SOW_APPROVED", project, {"revision": sow.revision_name})
    elif status_val == "in_review":
        project.stage = workflow.SOW_REVIEW
        workflow.log_action(db, user, "stage:SOW_REVIEW", project, {"revision": sow.revision_name})

    db.commit()
    db.refresh(sow)
    return sow


# ---------- BOQ / MTO Endpoints ----------
#: Keyword map used to place an extracted scope line under a trade.
_TRADE_HINTS: dict[str, tuple[str, ...]] = {
    "civil_arch": (
        "wall", "gypsum", "paint", "ceiling", "floor", "door", "partition",
        "dismantl", "casework", "cabinet", "concrete", "tile", "sealant", "curtain",
    ),
    "electrical": (
        "socket", "conduit", "cable", "wire", "panel", "breaker", "\bcb\b", "emt",
        "tagging", "commissioning", "earthing", "lighting", "ups", "thhn", "mcb",
    ),
    "low_current": (
        r"\bdata\b", "network", "cctv", "telecom", "fibre", "fiber", "router",
        "access control", "\blan\b",
    ),
    "plumbing": (
        "pipe", "drain", "valve", "water", "\bla\b", "lvac", "sink", "sewer", "ppr",
        "pex", "sanitary", "glasswasher",
    ),
    "hvac": (
        "duct", "\bfcu\b", "\bahu\b", "chiller", "\bair\b", "ventilat", "exhaust",
        "thermostat", "split unit", "balancing",
    ),
    "fire_protection": (
        "sprinkler", "fire alarm", "extinguisher", "smoke detector", "fm200",
    ),
    "tgm": ("gas detector", "\bh2\b", "c3h6", "c2h4", "tgm", "\bpds\b"),
}


#: A real scope line starts with one of these works (after bullets/numbers).
_SCOPE_VERB_RE = __import__("re").compile(
    r"^(supply|install|dismantl|relocat|remov|provid|test|testing|modif|construct|"
    r"replac|connect|extend|seal|paint|demolish|fabricat|upgrad|renew|shift|"
    r"re-?install|re-?locat|re-?work|cut|drill|mount|fix|run|lay|cable|wire|"
    r"terminat|commission|calibrat|program|integrat|abandon|patch|grind|coat|"
    r"excavat|core|break|restore|reinstate)",
    __import__("re").I,
)
#: PDF/Docx artefacts that are never scope items.
_ARTEFACT_RE = __import__("re").compile(
    r"(submittal|approver|approved\b|approval\]|project\s*:\s*\d|plan\s+north|"
    r"description please|comment please|university of science|page \d|"
    r"date\s*:|signature|revision\s+history|table of contents)",
    __import__("re").I,
)


def _clean_scope_line(raw: str) -> str | None:
    """Normalise a candidate line; None when it is not a scope item."""
    import re as _re

    text = " ".join(str(raw).split())
    text = text.lstrip("\u2022\u00b7-\u2013\u2014* ").strip()
    text = _re.sub(r"^\d+(\.\d+)*[\.\)]?\s*", "", text).strip()
    if len(text) < 14 or len(text) > 200:
        return None
    if _ARTEFACT_RE.search(text):
        return None
    letters = sum(ch.isalpha() for ch in text)
    if letters < len(text) * 0.6:
        return None
    if not _SCOPE_VERB_RE.match(text):
        return None
    return text


def _classify_scope_line(line: str) -> str | None:
    """Best-matching trade for a scope line (highest keyword score)."""
    import re as _re

    low = line.lower()
    best: tuple[int, str] | None = None
    for trade, hints in _TRADE_HINTS.items():
        score = sum(1 for h in hints if _re.search(h, low))
        if score and (best is None or score > best[0]):
            best = (score, trade)
    return best[1] if best else None


def _archive_scope_lines(project: Project, limit_projects: int = 6,
                         per_trade: int = 8) -> dict[str, list[dict]]:
    """Mine trade-wise scope lines from similar archived SOW documents."""
    import re as _re
    from pathlib import Path as _Path

    from ..services import archive_index as ai_idx
    from .projects import _archive_index_file, _derive_tracker_fields

    idx = ai_idx.load_index(_archive_index_file())
    derived = _derive_tracker_fields(project)
    hits = ai_idx.similar(
        idx, project.title, trades=derived.get("trades"),
        division=derived.get("phase"), exclude_pr=project.pr_number,
        limit=limit_projects,
    )
    out: dict[str, list[dict]] = {t: [] for t in _TRADE_HINTS}
    seen: set[str] = set()

    def _docx_lines(path) -> list[str]:
        try:
            from docx import Document

            return [p.text for p in Document(path).paragraphs]
        except Exception:  # noqa: BLE001
            return []

    def _pdf_lines(path) -> list[str]:
        """Text of the first few pages of an archived SOW PDF."""
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            chunks: list[str] = []
            for page in reader.pages[:6]:
                try:
                    chunks.append(page.extract_text() or "")
                except Exception:  # noqa: BLE001 - skip an unreadable page
                    continue
            return "\n".join(chunks).splitlines()
        except Exception:  # noqa: BLE001
            return []

    for hit in hits:
        folder = _Path(hit.get("path") or "")
        if not folder.is_dir() or "sow" not in (hit.get("doc_kinds") or []):
            continue
        # Archived SOWs are mostly PDF (PR9848-SOW-M-001.0.pdf, "PR-12245 Scope
        # of Work.pdf"); a few are Word. Read whichever exists.
        docs = [f for f in folder.rglob("*")
                if f.is_file()
                and f.suffix.lower() in (".pdf", ".docx")
                and _re.search(r"sow|scope of work", f.name, _re.I)
                and not f.name.startswith("~$")]
        docs.sort(key=lambda f: (f.suffix.lower() != ".pdf", len(f.name)))
        if not docs:
            continue
        raw_lines: list[str] = []
        for doc in docs[:2]:
            raw_lines.extend(_docx_lines(doc) if doc.suffix.lower() == ".docx" else _pdf_lines(doc))
            if raw_lines:
                break
        for raw in raw_lines:
            text = _clean_scope_line(raw)
            if not text:
                continue
            low = text.lower()
            key = low[:80]
            if key in seen:
                continue
            trade = _classify_scope_line(text)
            if not trade:
                continue
            seen.add(key)
            if len(out[trade]) < per_trade:
                out[trade].append({
                    "text": text,
                    "source_pr": hit.get("pr_number"),
                    "source_title": hit.get("title"),
                })
    return out


@router.get("/sow/suggestions")
def sow_suggestions(
    project_id: int,
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
):
    """Trade-wise scope suggestions for this project.

    Two sources, both for review - nothing is added automatically:
      * similar archived SOW documents (evidence of how the same work was
        scoped before), and
      * the project's own BOQ/MTO items already entered.
    """
    project = get_project_or_404(project_id, db)

    suggested = _archive_scope_lines(project)
    existing = _design_boq_items(project)
    existing_by_trade: dict[str, list[str]] = {}
    for it in existing:
        existing_by_trade.setdefault((it.trade or "general"), []).append(it.description)

    trades = []
    for trade, lines in suggested.items():
        if not lines:
            continue
        traded_existing = existing_by_trade.get(trade, [])
        already = {d.strip().lower() for d in traded_existing}
        fresh = [ln for ln in lines if ln["text"].strip().lower() not in already]
        if not fresh and not traded_existing:
            continue
        trades.append({
            "trade": trade,
            "suggestions": fresh,
            "existing_count": len(traded_existing),
        })
    return {
        "project": project.pr_number,
        "trades": trades,
        "total_suggestions": sum(len(t["suggestions"]) for t in trades),
    }


@router.get("/boq", response_model=list[BoqItemOut])
def list_boq_items(
    project_id: int,
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
):
    """List all BOQ/MTO line items for a project."""
    project = get_project_or_404(project_id, db)
    return project.boq_items


@router.post("/boq", response_model=BoqItemOut)
def add_boq_item(
    project_id: int,
    body: BoqItemCreate,
    db: Session = Depends(get_db),
    user: User = Depends(require_capability(CAP_BOQ_MANAGE)),
):
    """Add a structured BOQ line item."""
    project = get_project_or_404(project_id, db)
    total = float(body.quantity) * float(body.unit_rate)

    item = BoqMtoItem(
        project_id=project.id,
        trade=body.trade,
        item_code=body.item_code,
        description=body.description,
        unit=body.unit,
        quantity=body.quantity,
        unit_rate=body.unit_rate,
        total_rate=total,
        material_spec=body.material_spec,
        supplier_lead_time_days=body.supplier_lead_time_days,
        delivery_status="pending",
    )
    db.add(item)
    workflow.log_action(
        db,
        user,
        "boq:item_added",
        project,
        {"trade": body.trade, "item_code": body.item_code, "total": total},
    )
    db.commit()
    db.refresh(item)
    return item


@router.patch("/boq/{item_id}", response_model=BoqItemOut)
def update_boq_item(
    project_id: int,
    item_id: int,
    body: BoqItemUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(require_capability(CAP_BOQ_MANAGE)),
):
    """Update a BOQ line item."""
    project = get_project_or_404(project_id, db)
    item = db.get(BoqMtoItem, item_id)
    if not item or item.project_id != project.id:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="BOQ item not found"
        )

    if body.trade is not None:
        item.trade = body.trade
    if body.item_code is not None:
        item.item_code = body.item_code
    if body.description is not None:
        item.description = body.description
    if body.unit is not None:
        item.unit = body.unit
    if body.quantity is not None:
        item.quantity = body.quantity
    if body.unit_rate is not None:
        item.unit_rate = body.unit_rate
    if body.material_spec is not None:
        item.material_spec = body.material_spec
    if body.supplier_lead_time_days is not None:
        item.supplier_lead_time_days = body.supplier_lead_time_days
    if body.delivery_status is not None:
        item.delivery_status = body.delivery_status

    item.total_rate = float(item.quantity) * float(item.unit_rate)
    item.updated_at = datetime.now(timezone.utc)

    db.commit()
    db.refresh(item)
    return item


@router.delete("/boq/{item_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_boq_item(
    project_id: int,
    item_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_capability(CAP_BOQ_MANAGE)),
):
    """Delete a BOQ line item."""
    project = get_project_or_404(project_id, db)
    item = db.get(BoqMtoItem, item_id)
    if not item or item.project_id != project.id:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="BOQ item not found"
        )

    db.delete(item)
    workflow.log_action(
        db,
        user,
        "boq:item_deleted",
        project,
        {"trade": item.trade, "item_code": item.item_code},
    )
    db.commit()


@router.get("/boq/export")
def export_boq_excel(
    project_id: int,
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
):
    """Export the project BOQ/MTO line items to an Excel (.xlsx) workbook."""
    project = get_project_or_404(project_id, db)
    items = project.boq_items

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOQ & MTO Summary"

    # Styling definitions
    header_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="003366")
    bold_font = Font(name="Calibri", size=10, bold=True)
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    # Title block
    ws["A1"] = f"KAUST IN-HOUSE PROJECTS — BILL OF QUANTITIES & MATERIAL TAKE-OFF (BOQ/MTO)"
    ws["A1"].font = title_font
    ws["A2"] = f"PR: {project.pr_number} | Project: {project.title} | Location: {project.location or 'N/A'}"
    ws["A2"].font = bold_font
    ws["A3"] = f"Export Date: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A3"].font = regular_font

    headers = [
        "Item #",
        "Discipline / Trade",
        "Item Code",
        "Description & Specification",
        "Unit",
        "Qty",
        "Unit Rate (SAR)",
        "Total (SAR)",
        "Lead Time (Days)",
        "Status",
    ]

    row_num = 5
    for col_num, h_text in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_num = 6
    grand_total = 0.0
    for idx, itm in enumerate(items, 1):
        ws.cell(row=row_num, column=1, value=idx).font = regular_font
        ws.cell(row=row_num, column=2, value=itm.trade.replace("_", " ").title()).font = regular_font
        ws.cell(row=row_num, column=3, value=itm.item_code).font = regular_font
        ws.cell(row=row_num, column=4, value=f"{itm.description} {('(' + itm.material_spec + ')') if itm.material_spec else ''}").font = regular_font
        ws.cell(row=row_num, column=5, value=itm.unit).font = regular_font
        ws.cell(row=row_num, column=6, value=itm.quantity).font = regular_font
        ws.cell(row=row_num, column=7, value=itm.unit_rate).font = regular_font
        ws.cell(row=row_num, column=8, value=itm.total_rate).font = bold_font
        ws.cell(row=row_num, column=9, value=itm.supplier_lead_time_days or "N/A").font = regular_font
        ws.cell(row=row_num, column=10, value=itm.delivery_status.upper()).font = regular_font

        for c in range(1, 11):
            ws.cell(row=row_num, column=c).border = thin_border
        grand_total += itm.total_rate
        row_num += 1

    # Total row
    row_num += 1
    ws.cell(row=row_num, column=7, value="GRAND TOTAL (SAR):").font = bold_font
    tot_cell = ws.cell(row=row_num, column=8, value=grand_total)
    tot_cell.font = bold_font
    tot_cell.fill = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")

    # Column width formatting
    col_widths = {1: 8, 2: 18, 3: 14, 4: 42, 5: 8, 6: 10, 7: 15, 8: 18, 9: 16, 10: 14}
    for c, width in col_widths.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"BOQ_{project.pr_number}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------- Document Generation (docs/SOW_BOQ_MTO_PROMPT.md §3–§11) ----------

_DOCX_MEDIA = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_ZIP_MEDIA = "application/zip"


def _latest_sow(project: Project) -> SowRecord | None:
    return project.sow_records[-1] if project.sow_records else None


def _design_boq_items(project: Project) -> list[BoqMtoItem]:
    return [i for i in project.boq_items if i.mto_kind == "design"]


@router.post("/generate/{doc_type}")
def generate_document(
    project_id: int,
    doc_type: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """GENERATE SOW / BOQ / MTO / COMBINE PACKAGE (prompt §12).

    Deterministic generation from tagged data — SOW from the latest
    revision's trade_sections, BOQ from design BoqMtoItem rows, MTO from
    the trade_sections measurement lines. VAT/USD come from runtime
    settings. Capability: sow.manage for sow/package, boq.manage for
    boq/mto. ICR projects (no SOW stage) get 409 on sow/package.
    """
    project = get_project_or_404(project_id, db)

    if doc_type not in ("sow", "boq", "mto", "package"):
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="doc_type must be one of: sow, boq, mto, package",
        )
    needed = CAP_SOW_MANAGE if doc_type in ("sow", "package") else CAP_BOQ_MANAGE
    if needed not in effective_permissions(user):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=f"Missing permission: {needed}",
        )
    if project.disposition == "ICR" and doc_type in ("sow", "package"):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="SOW documents are not part of the ICR fast-track workflow.",
        )

    sow = _latest_sow(project)
    items = _design_boq_items(project)
    if sow is None and doc_type in ("sow", "package"):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Create a SOW revision (with trade_sections) before generating.",
        )

    if doc_type == "sow":
        path = docgen.generate_sow_docx(project, sow)
        filename = f"{project.pr_number} Scope of Work Draft.docx"
        media = _DOCX_MEDIA
    elif doc_type == "boq":
        path = docgen.generate_boq_xlsx(project, items, sow)
        filename = f"{project.pr_number} BOQ.xlsx"
        media = _XLSX_MEDIA
    elif doc_type == "mto":
        path = docgen.generate_mto_xlsx(project, sow)
        filename = f"{project.pr_number} MTO Design.xlsx"
        media = _XLSX_MEDIA
    else:
        path, _qa = docgen.generate_package(project, sow, items, user)
        filename = f"{project.pr_number}-Package.zip"
        media = _ZIP_MEDIA

    workflow.log_action(
        db, user, f"docgen:{doc_type}", project,
        {"file": filename, "revision": sow.revision_name if sow else None},
    )
    db.commit()
    return FileResponse(path, filename=filename, media_type=media)


@router.get("/qa-checklist")
def qa_checklist(
    project_id: int,
    db: Session = Depends(get_db),
    _user: User = Depends(require_capability(CAP_SOW_MANAGE)),
):
    """The §11 ten-point QA checklist. Any FAIL blocks package finalize."""
    project = get_project_or_404(project_id, db)
    return docgen.run_qa_checks(project, _latest_sow(project), _design_boq_items(project))
