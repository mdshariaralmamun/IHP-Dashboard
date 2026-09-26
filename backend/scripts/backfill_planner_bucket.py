"""Backfill projects.planner_bucket from the IHP planner tracker.

Writes ONLY the bucket column — no stage, disposition, or field changes —
so it is safe to run at any time (unlike a full planner import, which may
advance stages). Run once after the planner_bucket column is added, or any
time buckets look stale.

Usage:
    cd backend
    python -m scripts.backfill_planner_bucket [path-to-tracker.xlsx]
"""

from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl  # noqa: E402
from sqlalchemy import select  # noqa: E402

from app.core.config import get_settings  # noqa: E402
from app.db import SessionLocal  # noqa: E402
from app.models import Project  # noqa: E402
from app.services.tracker_files import latest_dated  # noqa: E402
from scripts.import_planner import parse_row  # noqa: E402


def _default_tracker() -> Path | None:
    """Always the newest dated IHP planner file — never a hardcoded date."""
    return latest_dated(get_settings().TRACKERS_DIR, "IHP- Construction Projects")


def backfill(path: Path, from_row: int = 10) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet_name = "Project tasks" if "Project tasks" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]

    db = SessionLocal()
    set_count = missing = 0
    distribution: Counter = Counter()
    try:
        for r in range(from_row, ws.max_row + 1):
            row = tuple(ws.cell(r, c).value for c in range(1, ws.max_column + 1))
            if not any(row):
                continue
            rec = parse_row(row)
            if rec.get("_skip") or not rec.get("pr_number"):
                continue
            project = db.scalar(
                select(Project).where(Project.pr_number == rec["pr_number"])
            )
            if project is None:
                missing += 1
                continue
            bucket = rec["bucket"] or None
            if project.planner_bucket != bucket:
                project.planner_bucket = bucket
                set_count += 1
            distribution[bucket or "(none)"] += 1
        db.commit()
    finally:
        db.close()
    return {"rows_seen": sum(distribution.values()), "buckets_set": set_count,
            "projects_not_in_db": missing, "distribution": dict(distribution)}


def main() -> int:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else _default_tracker()
    if path is None or not path.exists():
        print(f"ERROR: tracker not found: {path}")
        return 1
    result = backfill(path)
    print(f"Backfill from {path}")
    print(f"  rows seen       : {result['rows_seen']}")
    print(f"  buckets set     : {result['buckets_set']}")
    print(f"  not in DB       : {result['projects_not_in_db']}")
    print("  distribution    :")
    for bucket, count in sorted(result["distribution"].items()):
        print(f"    {bucket!r:>24}: {count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
