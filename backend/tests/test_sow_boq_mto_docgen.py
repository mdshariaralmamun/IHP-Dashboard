"""SOW + BOQ + MTO auto-generator tests (docs/SOW_BOQ_MTO_PROMPT.md).

Fixtures use the prompt's §8 worked example (PR-12630, socket relocation)
with the PR renamed PR-9005 so repeated runs stay unique. Everything is
asserted against the spec: §3.1 structure, §3.3 verbatim text blocks,
§4.4 BOQ sheets, §5.3 MTO sheet, §10 package tree, §11 QA checklist.
"""

import io
import itertools
import zipfile
from types import SimpleNamespace

import openpyxl
import pytest
from docx import Document as DocxDocument

from app.services.sow_boq_mto_docgen import (
    SOW_CIVIL_STD,
    SOW_GENERAL_STD,
    SOW_INTRO_STD,
    SOW_DOCS_STD,
    find_soffice,
)

PR = "PR-9005"
TITLE = "Relocation of 3-phase sockets"
LOCATION = "Building 6, Level-1, Greenhouse"

# The test DB is shared across the session, so each test seeds a unique PR.
_pr_seq = itertools.count()

_E = {"unit": "Lot", "qty": "1", "source_tag": "DOC", "source_doc": PR}
TRADE_SECTIONS = {
    "metadata": {
        "ear_number": "12547",
        "requester": "John Rahmer",
        "end_user": "Angelo Gallone",
        "funding": "OPEX",
        "wbs": "12380",
        "division": "Growth Chambers and Facilities",
        "estimated_cost": "USD 480 (excl. VAT)",
    },
    "drawings": [],  # none uploaded — QA check 6 must FAIL (§11)
    "trades": [
        {"name": "Civil / Architectural", "seen": True, "items": [
            {"description": "Gypsum wall mod, reworks, repairs, repainting, "
                            "fire-rated sealant", "source_doc": "EAR", **_E},
        ], "mto_items": [
            {"ref": "A.1", "description": "Gypsum board wall modification",
             "unit": "Lot", "qty": 1, "computation": "Lump sum"},
        ]},
        {"name": "Electrical", "seen": True, "items": [
            {"description": "Supply & install 1-inch EMT conduit + accessories", **_E},
            {"description": "Supply & install 1-inch flexible EMT conduit + accessories", **_E},
            {"description": "Relocate 3-phase socket UN3200-PR-HA-38,40,42 "
                            "from corridor to same room", **_E},
            {"description": "Testing & Commissioning + Tagging + Panel board schedule", **_E},
        ], "mto_items": [
            {"ref": "B.1", "description": "1-inch EMT conduit", "unit": "Lm",
             "qty": "TBC", "computation": "(from drawing)"},
            {"ref": "B.2", "description": "1-inch flexible EMT conduit", "unit": "Lm",
             "qty": "TBC", "computation": "(from drawing)"},
            {"ref": "B.3", "description": "Relocation of 3-phase socket", "unit": "Nos",
             "qty": 3, "computation": "3 sockets"},
            {"ref": "B.4", "description": "Cable pulling", "unit": "Lm",
             "qty": "TBC", "computation": "(from drawing)"},
            {"ref": "B.5", "description": "Testing & Commissioning", "unit": "Lot",
             "qty": 1, "computation": "Lump sum"},
        ]},
        {"name": "Low Current", "seen": False, "items": [], "mto_items": []},
        {"name": "Plumbing", "seen": False, "items": [], "mto_items": []},
        {"name": "HVAC", "seen": False, "items": [], "mto_items": []},
        {"name": "Fire Sprinkler", "seen": False, "items": [], "mto_items": []},
    ],
}

BOQ_LINES = [
    ("civil_architectural", "A.1",
     "Gypsum wall mod, reworks, repairs, repainting, fire-rated sealant"),
    ("electrical", "B.1", "Supply 1-inch EMT conduit with accessories"),
    ("electrical", "B.2", "Testing & Commissioning + Tagging + Panel schedule"),
]


@pytest.fixture()
def seeded_project(client, admin_headers):
    """A unique PR-9005-n with a Rev-0 SOW (§8 scope + MTO lines) and 3
    BOQ lines (§8.4). Returns SimpleNamespace(pid=..., pr=...)."""
    pr = f"{PR}-{next(_pr_seq)}"
    resp = client.post("/api/projects", headers=admin_headers, json={
        "pr_number": pr, "ear_number": "12547", "title": TITLE,
        "location": LOCATION, "funding_source": "OPEX",
    })
    assert resp.status_code in (200, 201), resp.text
    pid = resp.json()["id"]

    resp = client.post(f"/api/projects/{pid}/sow", headers=admin_headers,
                       json={"trade_sections": TRADE_SECTIONS})
    assert resp.status_code in (200, 201), resp.text

    for trade, code, desc in BOQ_LINES:
        resp = client.post(f"/api/projects/{pid}/boq", headers=admin_headers, json={
            "trade": trade, "item_code": code, "description": desc,
            "unit": "Lot", "quantity": 1, "unit_rate": 0,  # unpriced → TBC (§4.2)
        })
        assert resp.status_code in (200, 201), resp.text
    return SimpleNamespace(pid=pid, pr=pr)


def _docx_text(content: bytes) -> str:
    """All paragraph text + table cell text (tables are not in .paragraphs)."""
    doc = DocxDocument(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _sheet_text(ws) -> str:
    return " | ".join(str(c.value) for row in ws.iter_rows() for c in row
                      if c.value is not None)


class TestSowGeneration:
    def test_docx_structure_and_verbatim_blocks(self, client, admin_headers, seeded_project):
        resp = client.post(f"/api/projects/{seeded_project.pid}/generate/sow",
                           headers=admin_headers)
        assert resp.status_code == 200, resp.text
        assert "wordprocessingml" in resp.headers["content-type"]

        text = _docx_text(resp.content)
        # §3.3 verbatim (title/location substituted in the intro only)
        assert SOW_INTRO_STD.format(title=TITLE, location=LOCATION) in text
        assert SOW_CIVIL_STD in text
        assert SOW_GENERAL_STD in text
        assert SOW_DOCS_STD in text
        # §3.1 order — all six sections, "Not Seen" for the four empty trades
        for heading in ("2.1 Civil / Architectural", "2.2 Electrical",
                        "2.3 Low Current / Telecommunication", "2.4 Plumbing",
                        "2.5 HVAC", "2.6 Fire Sprinkler System"):
            assert heading in text, heading
        assert text.count("(Not Seen)") == 4
        # §8.3 electrical bullets present
        assert "UN3200-PR-HA-38,40,42" in text
        # §7 — tags preserved (traceability annex)
        assert "Annex A" in text and "Source Tag" in text

    def test_requires_a_sow_revision(self, client, admin_headers):
        resp = client.post("/api/projects", headers=admin_headers, json={
            "pr_number": f"PR-9006-{next(_pr_seq)}", "title": "No SOW yet",
        })
        assert resp.status_code in (200, 201)
        resp = client.post(f"/api/projects/{resp.json()['id']}/generate/sow",
                           headers=admin_headers)
        assert resp.status_code == 422

    def test_team_member_forbidden(self, client, role_headers, seeded_project):
        resp = client.post(f"/api/projects/{seeded_project.pid}/generate/sow",
                           headers=role_headers["member1"])
        assert resp.status_code == 403

    def test_unknown_doc_type_404(self, client, admin_headers, seeded_project):
        resp = client.post(f"/api/projects/{seeded_project.pid}/generate/nope",
                           headers=admin_headers)
        assert resp.status_code == 404


class TestBoqGeneration:
    def test_three_sheets_and_layout(self, client, admin_headers, seeded_project):
        resp = client.post(f"/api/projects/{seeded_project.pid}/generate/boq",
                           headers=admin_headers)
        assert resp.status_code == 200, resp.text
        assert "spreadsheetml" in resp.headers["content-type"]

        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        assert wb.sheetnames == ["Cover", "Bill of Quantity", "Take-off Trade"]

        cover = _sheet_text(wb["Cover"])
        for expected in (seeded_project.pr, "12547", TITLE, "John Rahmer",
                         "Angelo Gallone", "OPEX", "12380", "USD 480 (excl. VAT)"):
            assert expected in cover, expected

        ws = wb["Bill of Quantity"]
        headers = [ws.cell(row=6, column=c).value for c in range(1, 11)]
        assert headers == ["REF", "DESCRIPTION", "Unit", "QTY", "SUPPLY U.P.",
                           "SUPPLY TOTAL", "INSTALL U.P.", "INSTALL TOTAL",
                           "S.+I. UNIT", "S.+I. TOTAL"]
        text = _sheet_text(ws)
        # §4.2 — unpriced lines and totals are TBC; §4.1 groups A/B present
        assert "A.1" in text and "B.1" in text and "B.2" in text
        assert "TOTAL (SAR)" in text and "TOTAL (USD)" in text
        assert "VAT 15%" in text and "TOTAL + VAT (SAR)" in text
        assert text.count("TBC") >= 15
        # A–F group headers in canonical order
        order = [text.index(n) for n in
                 ("Civil / Architectural", "Electrical", "Plumbing", "HVAC",
                  "Low Current", "Fire Sprinkler")]
        assert order == sorted(order)

    def test_priced_lines_compute_totals(self, client, admin_headers, seeded_project):
        items = client.get(f"/api/projects/{seeded_project.pid}/boq",
                           headers=admin_headers).json()
        item_id = items[0]["id"]  # ids are global — never assume 1
        resp = client.patch(f"/api/projects/{seeded_project.pid}/boq/{item_id}",
                            headers=admin_headers, json={"unit_rate": 100.0})
        assert resp.status_code == 200, resp.text
        resp = client.post(f"/api/projects/{seeded_project.pid}/generate/boq",
                           headers=admin_headers)
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        # one priced line (1 × 100) → subtotal 100, VAT 15, total 115, USD @3.75
        values = [c.value for row in wb["Bill of Quantity"].iter_rows()
                  for c in row if c.value is not None]
        assert 100.0 in values          # SAR subtotal
        assert 115.0 in values          # SAR + VAT
        assert 15.0 in values           # VAT amount
        assert round(115.0 / 3.75, 2) in values  # USD total + VAT


class TestMtoGeneration:
    def test_mto_design_sheet(self, client, admin_headers, seeded_project):
        resp = client.post(f"/api/projects/{seeded_project.pid}/generate/mto",
                           headers=admin_headers)
        assert resp.status_code == 200, resp.text
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        assert wb.sheetnames == ["MTO Design"]
        ws = wb.active
        headers = [ws.cell(row=4, column=c).value for c in range(1, 8)]
        assert headers == ["Ref", "Description", "Unit", "QTY", "Details",
                           "Computation", "Total"]
        text = _sheet_text(ws)
        assert "3 sockets" in text          # §8.5 computation basis
        assert "(from drawing)" in text     # TBC measurement lines
        assert "Lump sum" in text
        assert "B.3" in text and "B.5" in text


class TestQaChecklist:
    def test_ten_checks_no_drawings_fails_6(self, client, admin_headers, seeded_project):
        resp = client.get(f"/api/projects/{seeded_project.pid}/qa-checklist",
                          headers=admin_headers)
        assert resp.status_code == 200
        checks = {c["n"]: c for c in resp.json()}
        assert len(checks) == 10
        assert checks[6]["passed"] is False   # no drawing register
        for n in (1, 2, 3, 4, 5, 7, 8, 9, 10):
            assert checks[n]["passed"] is True, (n, checks[n])


class TestPackage:
    def test_zip_tree_and_blocked_finalize(self, client, admin_headers, seeded_project):
        resp = client.post(f"/api/projects/{seeded_project.pid}/generate/package",
                           headers=admin_headers)
        assert resp.status_code == 200, resp.text
        assert "zip" in resp.headers["content-type"]

        pr = seeded_project.pr
        zf = zipfile.ZipFile(io.BytesIO(resp.content))
        names = zf.namelist()
        assert f"{pr}-Package/01_SOW/{pr} Scope of Work Draft.docx" in names
        assert f"{pr}-Package/02_BOQ/{pr} BOQ.xlsx" in names
        assert f"{pr}-Package/03_MTO/{pr} MTO Design.xlsx" in names
        assert f"{pr}-Package/04_Drawings/README.txt" in names
        assert f"{pr}-Package/CHANGELOG.md" in names
        if find_soffice():
            assert f"{pr}-Package/01_SOW/{pr} Scope of Work Draft.pdf" in names

        changelog = zf.read(f"{pr}-Package/CHANGELOG.md").decode("utf-8")
        # §11 — failing check blocks finalize
        assert "DRAFT — QA FAILED, finalize blocked" in changelog
        assert "FAIL" in changelog  # check 6 failure is reported
