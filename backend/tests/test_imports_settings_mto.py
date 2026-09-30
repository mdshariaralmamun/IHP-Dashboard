"""Tests for the admin import, settings, and MTO/budget routers.

Covers:
* /api/admin/settings      — read, override, clear, AI connection test
* /api/admin/import/*      — planner + O&M upload, mismatch report, resolve
* /api/mto/*               — master pricing CRUD, budget summary, generate-budget
"""

from __future__ import annotations

import io
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pytest

from app.api import imports as imports_api
from app.db import SessionLocal
from app.services import runtime_settings, tracker_sources
from app.models import BoqMtoItem, MasterPricing, Project


# ---------------------------------------------------------------------------
# xlsx fixtures
# ---------------------------------------------------------------------------


def _build_planner_xlsx() -> bytes:
    """One project row: PR-9001 in the MOM bucket."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Project tasks"
    # Rows 1-9 are the header block; data starts at row 10.
    # Column positions mirror the verified MS Project export layout
    # (1-based cell -> 0-based tuple index used by the parsers):
    #   col 3 = Name, col 5 = Bucket, col 8 = Finish, col 10 = Start,
    #   col 12 = % complete, col 26 = Building, col 28 = Requestor/PI
    ws.cell(10, 3, "9001 - N2 line installation")
    ws.cell(10, 5, "MOM")
    ws.cell(10, 8, "2026-10-01")
    ws.cell(10, 10, "2026-09-01")
    ws.cell(10, 12, 25)
    ws.cell(10, 26, "Building 3")
    ws.cell(10, 28, "Planner Assigned")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_om_xlsx() -> bytes:
    """One O&M row for PR-9001 with a conflicting location."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = " In House Projects"  # leading space is the real sheet name
    # Header row 1; data from row 2. col 1 = PR (int), col 4 = title,
    # col 6 = requestor, col 7 = location, col 8 = location details,
    # col 12 = status
    ws.cell(2, 1, 9001)
    ws.cell(2, 4, "N2 line installation")
    ws.cell(2, 6, "OM Requestor")
    ws.cell(2, 7, "Building 4")  # differs from planner's "Building 3"
    ws.cell(2, 8, "Level 2, Area 4")
    ws.cell(2, 12, "In House Projects")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture()
def planner_file():
    return {"file": ("planner.xlsx", _build_planner_xlsx(),
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


@pytest.fixture()
def om_file():
    return {"file": ("om.xlsx", _build_om_xlsx(),
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


@pytest.fixture()
def isolated_import_dir(tmp_path, monkeypatch):
    """Keep test uploads out of the real scratch + tracker stores.

    `IMPORT_DIR` is the scratch folder an upload lands in first; the store
    it is published to (`tracker_sources.upload_dir`) and the Planner's drop
    folder are redirected too, so each test sees only the files it uploaded.
    """
    monkeypatch.setattr(imports_api, "IMPORT_DIR", tmp_path / "imports")
    monkeypatch.setattr(tracker_sources, "upload_dir", lambda: tmp_path / "trackers")
    monkeypatch.setattr(tracker_sources, "configured_dir", lambda: tmp_path / "drop")
    return tmp_path


@pytest.fixture()
def isolated_settings_file(tmp_path, monkeypatch):
    """Point the settings override store at a temp file. Path resolution
    moved to services/runtime_settings (reads DATA_DIR per call), so patch
    the resolver instead of a module-level constant."""
    monkeypatch.setattr(
        runtime_settings, "_overrides_path", lambda: tmp_path / "settings.json"
    )


# ---------------------------------------------------------------------------
# /api/admin/settings
# ---------------------------------------------------------------------------


class TestSettings:
    def test_non_admin_cannot_read(self, client, role_headers, isolated_settings_file):
        resp = client.get("/api/admin/settings", headers=role_headers["member1"])
        assert resp.status_code == 403

    def test_read_defaults(self, client, admin_headers, isolated_settings_file):
        resp = client.get("/api/admin/settings", headers=admin_headers)
        assert resp.status_code == 200
        body = resp.json()
        assert body["ai_base_url"].startswith("http")
        assert "overrides" in body and body["overrides"] == {}
        assert "AI_BASE_URL" in body["env"]

    def test_patch_and_clear_override(self, client, admin_headers, isolated_settings_file):
        resp = client.patch(
            "/api/admin/settings",
            headers=admin_headers,
            json={"ai_chat_model": "test-model"},
        )
        assert resp.status_code == 200
        assert resp.json()["ai_chat_model"] == "test-model"
        overrides = resp.json()["overrides"]
        # The generic key is kept for backwards compatibility and the
        # per-provider key is what the provider layer reads first.
        assert overrides["AI_CHAT_MODEL"] == "test-model"
        assert set(overrides.values()) == {"test-model"}

        resp = client.patch(
            "/api/admin/settings",
            headers=admin_headers,
            json={"clear": list(overrides.keys())},
        )
        assert resp.status_code == 200
        assert resp.json()["overrides"] == {}

    def test_patch_rejects_bad_theme(self, client, admin_headers, isolated_settings_file):
        resp = client.patch(
            "/api/admin/settings", headers=admin_headers, json={"theme": "neon"},
        )
        assert resp.status_code == 400

    def test_ai_test_unreachable_is_ok_false(self, client, admin_headers, isolated_settings_file):
        # Pin the provider to local Ollama pointing at a closed port so the
        # test never depends on (or hits) any real hosted AI service, even
        # when a developer's .env configures a working provider + API key.
        # Port 1 refuses connections instantly, so no long timeout in tests.
        client.patch(
            "/api/admin/settings", headers=admin_headers,
            json={"ai_provider": "ollama", "ai_base_url": "http://127.0.0.1:1"},
        )
        resp = client.post("/api/admin/settings/ai/test", headers=admin_headers)
        assert resp.status_code == 200
        body = resp.json()
        assert body["ok"] is False
        assert "hint" in body


# ---------------------------------------------------------------------------
# /api/admin/import/*
# ---------------------------------------------------------------------------


class TestImports:
    def test_non_admin_cannot_upload(self, client, role_headers, planner_file):
        resp = client.post(
            "/api/admin/import/planner", files=planner_file,
            data={"dry_run": "true"}, headers=role_headers["member1"],
        )
        assert resp.status_code == 403

    def test_rejects_non_xlsx(self, client, admin_headers):
        resp = client.post(
            "/api/admin/import/planner",
            files={"file": ("plan.txt", b"not an xlsx", "text/plain")},
            data={"dry_run": "true"},
            headers=admin_headers,
        )
        assert resp.status_code == 400

    def test_planner_dry_run_imports_nothing(
        self, client, admin_headers, planner_file, isolated_import_dir,
    ):
        resp = client.post(
            "/api/admin/import/planner", files=planner_file,
            data={"dry_run": "true"}, headers=admin_headers,
        )
        assert resp.status_code == 200
        body = resp.json()
        assert body["dry_run"] is True
        assert body["rows_processed"] == 1

        db = SessionLocal()
        try:
            assert (
                db.query(Project).filter_by(pr_number="PR-9001").first() is None
            )
        finally:
            db.close()

    def test_planner_import_creates_project(
        self, client, admin_headers, planner_file, isolated_import_dir,
    ):
        resp = client.post(
            "/api/admin/import/planner", files=planner_file,
            data={"dry_run": "false"}, headers=admin_headers,
        )
        assert resp.status_code == 200
        assert resp.json()["rows_processed"] == 1

        db = SessionLocal()
        try:
            project = db.query(Project).filter_by(pr_number="PR-9001").first()
            assert project is not None
            # bucket "MOM" maps to (MOM_CONFIRMED, PROJECT)
            assert project.stage == "MOM_CONFIRMED"
            assert project.disposition == "PROJECT"
            # the raw planner bucket is persisted for the register filter
            assert project.planner_bucket == "MOM"
        finally:
            db.close()

    def test_om_upload_then_mismatch_report_and_resolve(
        self, client, admin_headers, planner_file, om_file, isolated_import_dir,
    ):
        # Upload planner (applies the import) then the O&M sheet.
        resp = client.post(
            "/api/admin/import/planner", files=planner_file,
            data={"dry_run": "false"}, headers=admin_headers,
        )
        assert resp.status_code == 200
        resp = client.post(
            "/api/admin/import/om", files=om_file, headers=admin_headers,
        )
        assert resp.status_code == 200
        assert resp.json()["rows_parsed"] == 1

        # Mismatch report should mention PR-9001 with at least the
        # location conflict (planner "Building 3" vs O&M "Building 4").
        resp = client.get("/api/admin/import/mismatches", headers=admin_headers)
        assert resp.status_code == 200
        body = resp.json()
        assert body["total"] >= 1
        item = next(i for i in body["items"] if i["pr_key"] == "PR-9001")
        assert item["mismatches"], "expected at least one conflict for PR-9001"
        fields = {m["field"] for m in item["mismatches"]}
        assert "location" in fields

        # Resolve the location conflict from the planner side.
        resp = client.post(
            "/api/admin/import/mismatches/resolve",
            headers=admin_headers,
            data={"pr_key": "PR-9001", "field": "location", "source": "planner"},
        )
        assert resp.status_code == 200
        assert resp.json()["ok"] is True

        db = SessionLocal()
        try:
            project = db.query(Project).filter_by(pr_number="PR-9001").first()
            assert project.location == "Building 3"
        finally:
            db.close()

    def test_resolve_rejects_bad_source(self, client, admin_headers, isolated_import_dir):
        resp = client.post(
            "/api/admin/import/mismatches/resolve",
            headers=admin_headers,
            data={"pr_key": "PR-9001", "field": "location", "source": "gut"},
        )
        assert resp.status_code == 400


_XLSX_MIME = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


class TestUploadPublishing:
    """An upload must become the tracker version the whole app resolves.

    Regression: uploads were written to a scratch folder nothing else read
    — and that a container rebuild wiped — so the O&M sheet had no effect
    at all, the Planner rewrote DB rows but the app kept pointing at (and
    naming) the previous tracker file.
    """

    def _upload_planner(
        self, client, admin_headers, name="planner.xlsx", dry_run="false",
    ):
        return client.post(
            "/api/admin/import/planner",
            files={"file": (name, _build_planner_xlsx(), _XLSX_MIME)},
            data={"dry_run": dry_run},
            headers=admin_headers,
        )

    def test_planner_upload_is_published_and_resolved(
        self, client, admin_headers, isolated_import_dir,
    ):
        resp = self._upload_planner(client, admin_headers)
        assert resp.status_code == 200
        body = resp.json()
        assert body["published"] is True

        stored = tracker_sources.upload_dir() / body["active_file"]
        assert stored.is_file(), "the upload must land in the tracker store"
        assert tracker_sources.planner_path() == stored
        assert body["active_source"] == "upload"

    def test_undated_upload_is_stamped_with_todays_date(
        self, client, admin_headers, isolated_import_dir,
    ):
        stamp = datetime.now().strftime("%d%m%Y")
        resp = self._upload_planner(
            client, admin_headers, name="IHP- Construction Projects.xlsx",
        )
        assert resp.status_code == 200
        assert resp.json()["active_file"] == (
            f"IHP- Construction Projects_{stamp}.xlsx"
        )

    def test_a_misnamed_upload_is_stored_under_the_tracker_name(
        self, client, admin_headers, isolated_import_dir,
    ):
        """Resolution matches the tracker's filename prefix, so a file that
        arrives as "planner.xlsx" must be stored under the canonical name —
        otherwise it would sit in the store and never be resolved."""
        stamp = datetime.now().strftime("%d%m%Y")
        resp = self._upload_planner(client, admin_headers, name="planner.xlsx")
        assert resp.status_code == 200
        assert resp.json()["active_file"] == (
            f"IHP- Construction Projects_{stamp}.xlsx"
        )

    def test_a_misnamed_upload_keeps_the_date_in_its_name(
        self, client, admin_headers, isolated_import_dir,
    ):
        resp = self._upload_planner(
            client, admin_headers,
            name="Copy of IHP- Construction Projects_26092026.xlsx",
        )
        assert resp.status_code == 200
        assert resp.json()["active_file"] == (
            "IHP- Construction Projects_26092026.xlsx"
        )

    def test_upload_beats_an_older_file_in_the_drop_folder(
        self, client, admin_headers, isolated_import_dir,
    ):
        drop = isolated_import_dir / "drop"
        drop.mkdir(parents=True, exist_ok=True)
        (drop / "IHP- Construction Projects_07092026.xlsx").write_bytes(
            _build_planner_xlsx()
        )
        resp = self._upload_planner(
            client, admin_headers,
            name="IHP- Construction Projects_30092026.xlsx",
        )
        assert resp.status_code == 200
        assert tracker_sources.planner_path().name == (
            "IHP- Construction Projects_30092026.xlsx"
        )
        assert resp.json()["active_source"] == "upload"

    def test_an_older_upload_never_rolls_the_app_backwards(
        self, client, admin_headers, isolated_import_dir,
    ):
        """Priority is the DATE, not the folder: uploading an older export
        must not make the app serve it over a newer file already present."""
        resp = self._upload_planner(
            client, admin_headers,
            name="IHP- Construction Projects_07092026.xlsx",
        )
        assert resp.status_code == 200
        drop = isolated_import_dir / "drop"
        drop.mkdir(parents=True, exist_ok=True)
        newer = drop / "IHP- Construction Projects_26092026.xlsx"
        newer.write_bytes(_build_planner_xlsx())
        assert tracker_sources.planner_path() == newer
        assert tracker_sources.status()["planner_source"] == "folder"

    def test_dry_run_stages_without_publishing(
        self, client, admin_headers, isolated_import_dir,
    ):
        resp = self._upload_planner(client, admin_headers, dry_run="true")
        assert resp.status_code == 200
        assert resp.json()["published"] is False
        assert tracker_sources.planner_path() is None

    def test_om_upload_becomes_the_resolved_om_sheet(
        self, client, admin_headers, om_file, isolated_import_dir,
    ):
        resp = client.post(
            "/api/admin/import/om", files=om_file, headers=admin_headers,
        )
        assert resp.status_code == 200
        body = resp.json()
        assert body["published"] is True

        stored = tracker_sources.upload_dir() / body["active_file"]
        assert stored.is_file()
        assert tracker_sources.om_path() == stored
        assert tracker_sources.status()["om_source"] == "upload"

    def test_rejected_om_upload_is_never_published(
        self, client, admin_headers, isolated_import_dir,
    ):
        """A workbook the parser rejects must not become the live sheet."""
        wb = openpyxl.Workbook()
        wb.active.title = "Wrong Sheet"
        buf = io.BytesIO()
        wb.save(buf)
        resp = client.post(
            "/api/admin/import/om",
            files={"file": ("om.xlsx", buf.getvalue(), _XLSX_MIME)},
            headers=admin_headers,
        )
        assert resp.status_code == 400
        assert tracker_sources.om_path() is None
        assert not list(tracker_sources.upload_dir().glob("*.xlsx"))

    def test_reuploading_the_same_version_replaces_it(
        self, client, admin_headers, isolated_import_dir,
    ):
        first = self._upload_planner(client, admin_headers)
        second = self._upload_planner(client, admin_headers)
        assert first.json()["active_file"] == second.json()["active_file"]
        assert len(list(tracker_sources.upload_dir().glob("*.xlsx"))) == 1

    def test_sources_endpoint_reports_the_live_version(
        self, client, admin_headers, isolated_import_dir,
    ):
        self._upload_planner(client, admin_headers)
        resp = client.get("/api/admin/import/sources", headers=admin_headers)
        assert resp.status_code == 200
        body = resp.json()
        assert body["planner_source"] == "upload"
        assert body["upload_dir"] == str(tracker_sources.upload_dir())
        assert body["configured_dir"] == str(isolated_import_dir / "drop")


# ---------------------------------------------------------------------------
# /api/mto/*
# ---------------------------------------------------------------------------


_MTO_PR_SEQ = iter(range(9500, 9999))


@pytest.fixture()
def mto_project(client, admin_headers):
    pr_number = f"PR-{next(_MTO_PR_SEQ)}"
    resp = client.post(
        "/api/projects",
        headers=admin_headers,
        json={"pr_number": pr_number, "title": "MTO budget test project"},
    )
    assert resp.status_code in (200, 201), resp.text
    return resp.json()["id"]


class TestMtoPricing:
    def test_unknown_item_404(self, client, admin_headers):
        resp = client.get("/api/mto/pricing/NO-SUCH-ITEM", headers=admin_headers)
        assert resp.status_code == 404

    def test_upsert_persists_and_updates(self, client, admin_headers):
        payload = {
            "item_code": "PIPE-CU-05",
            "description": "Copper pipe 1/2 in",
            "trade": "plumbing",
            "unit": "L.M.",
            "base_unit_rate": 125.0,
        }
        resp = client.post("/api/mto/pricing", headers=admin_headers, json=payload)
        assert resp.status_code == 200
        assert resp.json()["base_unit_rate"] == 125.0

        # The write must survive the request (commit, not just flush).
        resp = client.get("/api/mto/pricing/PIPE-CU-05", headers=admin_headers)
        assert resp.status_code == 200

        # Upsert with a new rate -> single row, updated rate.
        payload["base_unit_rate"] = 130.0
        resp = client.post("/api/mto/pricing", headers=admin_headers, json=payload)
        assert resp.status_code == 200
        assert resp.json()["base_unit_rate"] == 130.0

        db = SessionLocal()
        try:
            rows = db.query(MasterPricing).filter(
                MasterPricing.item_code == "PIPE-CU-05"
            ).all()
            assert len(rows) == 1
            assert rows[0].base_unit_rate == 130.0
        finally:
            db.close()

    def test_list_filters_by_trade(self, client, admin_headers):
        client.post("/api/mto/pricing", headers=admin_headers, json={
            "item_code": "WIRE-2.5", "description": "Wire", "trade": "electrical",
            "base_unit_rate": 5.0,
        })
        resp = client.get(
            "/api/mto/pricing", headers=admin_headers, params={"trade": "electrical"},
        )
        assert resp.status_code == 200
        codes = [p["item_code"] for p in resp.json()]
        assert codes == ["WIRE-2.5"]

    def test_non_admin_cannot_write(self, client, role_headers):
        resp = client.post("/api/mto/pricing", headers=role_headers["member1"], json={
            "item_code": "X", "description": "x", "base_unit_rate": 1.0,
        })
        assert resp.status_code == 403


class TestMtoBudget:
    def test_budget_summary_auto_creates(self, client, admin_headers, mto_project):
        resp = client.get(
            f"/api/mto/project/{mto_project}/budget-summary", headers=admin_headers,
        )
        assert resp.status_code == 200
        body = resp.json()
        assert body["project_id"] == mto_project
        assert body["stage_at_snapshot"] == "INTAKE"

    def test_generate_budget_uses_master_pricing(
        self, client, admin_headers, mto_project,
    ):
        db = SessionLocal()
        try:
            db.add_all([
                MasterPricing(
                    item_code="GB-PIPE", description="Pipe", trade="plumbing",
                    unit="L.M.", base_unit_rate=125.0,
                ),
                BoqMtoItem(
                    project_id=mto_project, mto_kind="design", trade="plumbing",
                    item_code="GB-PIPE", description="Pipe run", unit="L.M.",
                    quantity=2.0, unit_rate=1.0,  # stale rate; master wins
                ),
                BoqMtoItem(
                    project_id=mto_project, mto_kind="design", trade="plumbing",
                    item_code="GB-VALVE", description="Valve", unit="EA",
                    quantity=1.0, unit_rate=80.0,  # no master row; item rate wins
                ),
            ])
            db.commit()
        finally:
            db.close()

        resp = client.post(
            f"/api/mto/project/{mto_project}/generate-budget", headers=admin_headers,
        )
        assert resp.status_code == 200, resp.text
        body = resp.json()
        summary = body["summary"]
        # 2 x 125 (master pricing) + 1 x 80 (stored rate) = 330
        assert summary["subtotal_sar"] == pytest.approx(330.0)
        assert summary["vat_sar"] == pytest.approx(49.5)
        assert summary["total_sar"] == pytest.approx(379.5)
        assert summary["total_usd"] == pytest.approx(round(379.5 / 3.75, 2))
        sources = {i["item_code"]: i["pricing_source"] for i in body["items_detail"]}
        assert sources["GB-PIPE"] == "master_pricing"
        assert sources["GB-VALVE"] == "item_stored"

    def test_generate_budget_no_items_400(self, client, admin_headers):
        resp = client.post(
            "/api/projects", headers=admin_headers,
            json={"pr_number": f"PR-{next(_MTO_PR_SEQ)}", "title": "Empty MTO project"},
        )
        assert resp.status_code in (200, 201), resp.text
        project_id = resp.json()["id"]
        resp = client.post(
            f"/api/mto/project/{project_id}/generate-budget", headers=admin_headers,
        )
        assert resp.status_code == 400

    def test_stage_budget_upsert(self, client, admin_headers, mto_project):
        resp = client.post(
            f"/api/mto/project/{mto_project}/stage-budget",
            headers=admin_headers, params={"stage": "MTO", "budget_sar": 5000.0},
        )
        assert resp.status_code == 200
        assert resp.json()["budget_sar"] == 5000.0

        # Upsert same stage with a new budget -> still one row
        resp = client.post(
            f"/api/mto/project/{mto_project}/stage-budget",
            headers=admin_headers, params={"stage": "MTO", "budget_sar": 6000.0},
        )
        assert resp.status_code == 200
        assert resp.json()["budget_sar"] == 6000.0

        resp = client.get(
            f"/api/mto/project/{mto_project}/stage-budget", headers=admin_headers,
        )
        rows = resp.json()
        assert len(rows) == 1
        assert rows[0]["budget_sar"] == 6000.0
