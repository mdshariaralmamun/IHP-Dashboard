"""KAUST MOM template generation: details payload, revision no., document code."""

from docx import Document

DETAILS = {
    "meeting_title": "PR-70001 kickoff — scope confirmation",
    "meeting_location": "Bldg 5, Level 3",
    "meeting_number": "03",
    "meeting_date": "2026-09-10",
    "meeting_time": "10:00",
    "attendees": [
        {
            "name": "Nikolay Gorshkov",
            "title": "Requester",
            "email": "nikolay.gorshkov@kaust.edu.sa",
        },
        {"name": "Mohammed Al Mamun", "title": "IHP Engineer", "email": "m@kaust.edu.sa"},
    ],
    "agenda": [
        {"scope": "Confirm chilled water tie-in point", "action": "Civil", "etc": "Site visit", "trade": "Plumbing"},
        {"scope": "Utility matrix review", "action": "MEP", "etc": "", "trade": "Electrical"},
    ],
}


def _create_project(client, admin_headers, pr):
    resp = client.post(
        "/api/projects",
        json={
            "pr_number": pr,
            "title": "Chilled water tie-in",
            "location": "Bldg 5",
            "pi_name": "Nikolay Gorshkov",
            "pi_email": "nikolay.gorshkov@kaust.edu.sa",
            "funding_source": "ASEPC",
        },
        headers=admin_headers,
    )
    assert resp.status_code == 201, resp.text
    return resp.json()["id"]


def _docx_text(path):
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_generate_with_details(client, admin_headers, data_dir):
    pid = _create_project(client, admin_headers, "PR-70001")
    resp = client.post(
        f"/api/projects/{pid}/mom/generate", json=DETAILS, headers=admin_headers
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["details"]["meeting_number"] == "03"

    text = _docx_text(data_dir / "projects" / "PR-70001" / "mom" / "MOM_PR-70001_v1.docx")
    assert "KING ABDULLAH UNIVERSITY OF SCIENCE AND TECHNOLOGY" in text
    assert "FACILITIES MANAGEMENT" in text
    assert "KA-HBRP-G-MOM" in text
    assert "Revision No.: 00" in text
    assert "PR-70001 kickoff — scope confirmation" in text
    assert "Bldg 5, Level 3" in text
    assert "2026-09-10" in text
    assert "Mohammed Al Mamun" in text
    assert "Confirm chilled water tie-in point" in text
    # House order: Civil/Architectural -> Plumbing -> HVAC -> Electrical ->
    # General, regardless of the order the items were entered in. Each row also
    # starts with its trade name so the minute reads trade by trade.
    assert text.index("Plumbing") < text.index("Electrical")
    assert text.index("Confirm chilled water tie-in point") < text.index("Utility matrix review")


def test_regenerate_keeps_details_and_bumps_revision(client, admin_headers, data_dir):
    pid = _create_project(client, admin_headers, "PR-70002")
    r1 = client.post(f"/api/projects/{pid}/mom/generate", json=DETAILS, headers=admin_headers)
    assert r1.status_code == 200, r1.text
    r2 = client.post(f"/api/projects/{pid}/mom/generate", headers=admin_headers)
    assert r2.status_code == 200, r2.text
    assert r2.json()["details"]["meeting_title"] == DETAILS["meeting_title"]

    # An explicit empty payload must also keep the stored details.
    r3 = client.post(f"/api/projects/{pid}/mom/generate", json={}, headers=admin_headers)
    assert r3.status_code == 200, r3.text
    assert r3.json()["details"]["meeting_title"] == DETAILS["meeting_title"]

    text = _docx_text(data_dir / "projects" / "PR-70002" / "mom" / "MOM_PR-70002_v2.docx")
    assert "Revision No.: 01" in text


def test_generate_defaults(client, admin_headers, data_dir):
    pid = _create_project(client, admin_headers, "PR-70003")
    resp = client.post(f"/api/projects/{pid}/mom/generate", headers=admin_headers)
    assert resp.status_code == 200, resp.text

    text = _docx_text(data_dir / "projects" / "PR-70003" / "mom" / "MOM_PR-70003_v1.docx")
    assert "Revision No.: 00" in text
    assert "PR-70003 — Chilled water tie-in" in text  # default meeting title
    assert "Nikolay Gorshkov" in text  # default PI attendee row
