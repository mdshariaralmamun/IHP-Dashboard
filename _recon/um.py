# -*- coding: utf-8 -*-
import openpyxl, sys
from openpyxl.utils import get_column_letter
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
ws = wb['Utility Matrix']
for r in (4,5):
    vals=[]
    for c in range(1,47):
        v = ws.cell(row=r,column=c).value
        if v is not None:
            vals.append(f"{get_column_letter(c)}{r}={str(v).strip()}")
    print("ROW",r,":"," | ".join(vals))
