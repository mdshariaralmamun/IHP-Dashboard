# -*- coding: utf-8 -*-
import sys, zipfile, re, io
import docx
from docx.shared import Pt

path = sys.argv[1]
out = io.StringIO()
def w(s=""):
    out.write(str(s) + "\n")

d = docx.Document(path)
w("FILE: " + path)
w("="*100)

# ---------- OUTLINE ----------
w("\n### OUTLINE (paragraphs with style Heading 1/2/3 or Title, in order) ###")
counts = {}
for p in d.paragraphs:
    st = p.style.name if p.style else ""
    if st.startswith("Heading") or st == "Title":
        counts[st] = counts.get(st,0)+1
        w(f"[{st}] {p.text}")
w("\nHeading style counts: " + str(counts))

# ---------- ALL PARAGRAPHS (compact, with styles for non-empty) ----------
w("\n### ALL NON-EMPTY PARAGRAPHS WITH STYLE (first 200) ###")
n=0
for i,p in enumerate(d.paragraphs):
    t = p.text.strip()
    if not t: continue
    n+=1
    if n>200:
        w("... (truncated)")
        break
    w(f"{i:4d} <{p.style.name}> {t}")

# ---------- TABLES ----------
w("\n### TABLES ###")
w("total tables: %d" % len(d.tables))
for ti,t in enumerate(d.tables):
    w("\n--- TABLE %d : rows=%d cols=%d ---" % (ti, len(t.rows), len(t.columns)))
    limit = min(len(t.rows), 4)  # header + 3 sample
    for ri in range(limit):
        cells = [c.text.replace("\n"," / ").strip() for c in t.rows[ri].cells]
        w(f"  R{ri}: {cells}")
    if len(t.rows) > 4:
        w(f"  ... ({len(t.rows)-4} more rows)")

sys.stdout.buffer.write(out.getvalue().encode("utf-8"))
