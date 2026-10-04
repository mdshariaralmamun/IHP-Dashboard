# -*- coding: utf-8 -*-
import sys, io
import openpyxl
path = sys.argv[1]
nrows = int(sys.argv[2]) if len(sys.argv)>2 else 18
outpath = sys.argv[3] if len(sys.argv)>3 else None
out = io.StringIO()
def w(s=""): out.write(str(s)+"\n")
wb = openpyxl.load_workbook(path, data_only=False)
w("FILE: "+path)
w("sheets: "+repr(wb.sheetnames))
for ws in wb.worksheets:
    w("\n"+"="*90)
    w(f"SHEET: {ws.title!r}  dims={ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column}")
    w("merged: "+repr([str(m) for m in ws.merged_cells.ranges][:40]))
    w("freeze: %s  auto_filter: %s" % (ws.freeze_panes, ws.auto_filter.ref))
    for r in range(1, min(ws.max_row, nrows)+1):
        vals=[]
        for c in range(1, min(ws.max_column, 20)+1):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                s = str(v).replace("\n"," / ")
                if len(s)>60: s = s[:57]+"..."
                vals.append(f"{openpyxl.utils.get_column_letter(c)}{r}={s}")
        if vals:
            w("  "+" | ".join(vals))
data = out.getvalue().encode("utf-8")
if outpath:
    open(outpath,"wb").write(data); print("wrote",outpath,len(data))
else:
    sys.stdout.buffer.write(data)
